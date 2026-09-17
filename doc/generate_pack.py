from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
import json, zipfile, textwrap, re, shutil, os

BASE = Path('/mnt/data/myanmar_logistics_docs')
ASSIGNMENT = Path('/mnt/data/README-From-assignment.txt').read_text()

# Decision log: 30 rounds x 10 questions. Choices preserve the grill structure.
# final notes capture user overrides and later refinements.
Q=[]
def q(n, question, A, B, C, D, rec, final=None, note=''):
    Q.append(dict(n=n, question=question, choices={'A':A,'B':B,'C':C,'D':D}, recommendation=rec,
                  final=final or rec, note=note))

# Round 1: Scope & Architecture
q(1,'What exactly are you trying to deliver?','A complete commercial logistics platform','A polished prototype demonstrating the required scenario','Mostly UI screens with no real backend','A technical backend with a simple UI','B')
q(2,'What should be the single most important demo story?','Driver uploads many logistics documents','Trader creates shipment → Driver moves → Gate closes/appears → Trader gets alerted → Shipment delivered','Admin views complicated analytics','Users chat with an AI logistics assistant','B')
q(3,'What technology stack should you standardize on?','Next.js + Tailwind + Supabase','React + custom Express backend + PostgreSQL','Flutter + Django + MySQL','Separate technologies chosen independently by each member','A','CUSTOM','Final stack: React/Vite + Tailwind frontend, Python/FastAPI backend, Supabase Auth/Postgres/Storage/Realtime; OpenRouteService for routing, IndexedDB for offline.')
q(4,'How should you implement the three roles?','Three totally separate applications','One app with authentication and role-based dashboards','No authentication; three different URLs','Hard-code different screens during presentation','B')
q(5,'How realistic should authentication be?','Build enterprise-grade permissions','Supabase authentication with role stored in a profile table and backend/RLS enforcement','Fake login buttons only','Skip login entirely','B')
q(6,'What should your shipment status model be?','Only Active / Completed','Requested → Assigned → Picked Up → In Transit → Checkpoint → Customs → Delivered','20+ detailed logistics states','Let Drivers type arbitrary status text','B','B','Later refined with CANCELLED and separate delay_status.')
q(7,'How should you simulate GPS?','Integrate only a real Driver phone GPS','Use predetermined coordinates and move the truck marker through them','Build your own GPS tracking service','Use only text locations','B','B','Use OpenRouteService for route geometry. Later upgraded to support both real production mobile GPS and demo simulation.')
q(8,'How should the Myanmar border/gate feature work?','Admin changes a database gate status and affected shipments receive an alert','Admin manually messages each Trader','Gate status is only decorative','Pull live Myanmar border-government data','A','CUSTOM','Driver can report existing or new gates/checkpoints; reports are evidence. Admin remains authoritative for manual overrides, while later design allows multi-driver consensus to auto-confirm/resolve routine incidents.')
q(9,'How much offline functionality should you actually build?','Full production offline synchronization','Simulate/implement offline mode: cache Driver updates locally and sync when connection returns','Mention offline support only in slides','Ignore it','B')
q(10,'With five people, what should you deliberately NOT build first?','AI chatbot, payments, OCR, heavy analytics and other non-core extras','Admin role','Driver interface','Gate alerts','A','A','Multiple/alternative routes were allowed later, but only after the core flow; learned routes are a differentiator, not a prerequisite.')

# Round 2: DB + gates + routes
q(11,'Who has final manual authority over official gate status?','Any Driver','Admin','Trader','Whoever reported it first','B','B','Routine consensus may auto-confirm later; Admin can manually override.')
q(12,'What should happen when a Driver reports a blocked checkpoint?','Immediately mark the whole route closed','Create a pending incident/report requiring verification or corroboration','Ignore it until another Driver reports it','Send only a chat message','B')
q(13,'What gate statuses should you support?','Open / Closed only','Open / Congested / Blocked / Closed','15 different statuses','Free text','B')
q(14,'What should Drivers submit with a gate report?','Status only','Status + location + timestamp + note + optional photo','Photo only','Long written report','B')
q(15,'How should shipments relate to routes?','Every shipment has one active route','Every shipment simultaneously follows several routes','Shipments do not store routes','Driver chooses randomly','A','A','Route history and alternatives are stored separately; active_route_id represents current route.')
q(16,'Who should approve a reroute after a gate closure?','Trader','Driver alone','Admin','Automatically reroute every shipment','C','B','Superseded later: Driver owns operational routing and can reroute their assigned shipment without Admin approval; route change is logged and visible.')
q(17,'How should OpenRouteService be used?','Make it the entire logistics backend','Use it to calculate/display route geometry while your app owns shipment state','Store everything inside ORS','Do not maintain route data yourself','B')
q(18,'What happens if OpenRouteService fails during the presentation?','Demo fails','Use saved/cached route coordinates or GeoJSON fallback','Restart the presentation','Remove the map completely','B')
q(19,'What should happen when the Driver is offline and reports a gate problem?','Report disappears','Save locally as UNSYNCED and upload when internet returns','Driver cannot use the application','Immediately notify Admin somehow','B')
q(20,'How far should multiple-route functionality go for the core prototype?','Full automatic dynamic route optimization','Primary route + alternative/Driver reroute with history','No route concept at all','AI dynamically invents routes','B','B','Later expanded with observed/unmapped routes and route recommendations after core requirements are stable.')

# Round 3: Database + API ownership
q(21,'What should be the core Supabase tables?','users, shipments only','profiles, shipments, routes, gates, gate_reports, shipment_events, alerts, documents','One giant logistics_data table','Separate database for every role','B','B','Later expanded with route_points, route_traversals and incidents. Gates are hybrid: predefined known gates plus temporary Driver-discovered checkpoints.')
q(22,'Where should user roles be stored?','React local storage only','profiles table linked to Supabase Auth','Hardcoded in frontend files','Browser cookies only','B')
q(23,'What should the shipments table contain?','Only shipment name','Trader, Driver, route, cargo, origin/destination, current state and timestamps','Every GPS coordinate ever recorded','Gate reports and notifications inside the same row','B')
q(24,'Where should shipment history be stored?','Overwrite shipment.status and keep no history','Separate shipment_events table','One huge text string','Store history in React','B')
q(25,'How should GPS location be represented?','Only one current latitude/longitude','Current position plus GPS/history events or batches','Driver types city name','Screenshots of maps','B')
q(26,'Where should Driver-uploaded photos be stored?','Huge binary values in PostgreSQL','Supabase Storage with paths/metadata in documents or reports','Driver phone only','React source folder','B')
q(27,'Where should important business logic live?','React only','FastAPI/backend services','CSS','Browser local storage','B')
q(28,'Should React directly access Supabase?','Never','Yes for safe auth/realtime/read operations under RLS; privileged actions go through FastAPI','React should have the server secret key','Everything must go through Python, including simple subscriptions','B')
q(29,'How should alerts work?','Hardcoded popup','Store alerts in DB and use Supabase Realtime to update relevant UIs','Send emails only','Refresh browser every minute','B')
q(30,'How should your five people divide the actual coding?','Everyone codes everything','Divide by ownership against one shared schema/API contract','Each builds a separate prototype','One person codes while four prepare slides','B')

# Round 4: screens + dynamic gates
q(31,'What should the Admin see immediately after login?','Blank dashboard with menu links','Shipment summary + active gate issues + pending Driver reports + alerts','Only user management','Raw database tables','B')
q(32,'What should happen when Admin opens a pending Driver gate report?','Only show its text','Show location, photo, Driver, shipment, timestamp, note and review controls','Automatically approve it','Delete it after reading','B')
q(33,'Should Admin be allowed to create a gate without a Driver report?','Yes','No','Only after two Driver reports','Only Traders can create gates','A')
q(34,'What should the Trader dashboard prioritize?','All Drivers in the system','Trader’s own shipments, current status, map, timeline and alerts','System administration','Gate approval','B')
q(35,'What should a Trader see on shipment detail?','Tracking number only','Cargo + Driver + route/map + timeline + current status + related alerts','Admin controls','Backend logs','B')
q(36,'What should the Driver home screen emphasize?','Analytics','Currently assigned shipment and large operational action buttons','All company shipments','User management','B')
q(37,'How should a Driver report a completely new checkpoint?','Force selection of an existing gate','Report New Checkpoint captures location and asks for status, note and optional photo','Type coordinates manually','Call Admin instead','B')
q(38,'What should happen when multiple Drivers report approximately the same temporary gate?','Always create separate gates','Show possible nearby match and allow association/merge into one gate/incident','Delete later reports','Automatically ban duplicate reports','B')
q(39,'What should happen after a temporary checkpoint disappears?','Delete everything','Mark it resolved/archived and stop active warnings while preserving history','Leave it permanently blocked','Driver deletes it directly','B')
q(40,'How many screens should your team build before polishing?','20+ screens','Minimum screens needed for one complete demo journey','Every possible logistics screen','Build every role simultaneously without prioritizing','B')

# Round 5: maps + demo
q(41,'What should the Admin map show?','Every user home address','Active trucks, routes, known gates, temporary checkpoints and incidents','Only Myanmar base map','Only one truck','B')
q(42,'What should the Trader map show?','Every truck','Only their shipment, route, truck position and relevant disruptions','Every checkpoint in Myanmar','Driver private GPS history','B')
q(43,'What should the Driver map show?','All customer shipments','Current location, assigned/available routes, destination and relevant checkpoints','Only destination','Admin analytics','B')
q(44,'How should simulated truck movement work during the presentation?','Wait for real GPS movement','Use predefined route coordinates with controlled manual/step simulation','Move icon randomly','Let ORS decide every few seconds','B')
q(45,'What should be the most impressive moment in the live demo?','Logging in','Driver discovers checkpoint → reports → incident confirmed → Trader map/alert updates','Opening database','Showing source code','B')
q(46,'Who should perform the live demo?','All 5 constantly switch laptops','1–2 operators while others explain specific sections','Only backend developer','Everyone clicks simultaneously','B')
q(47,'How should you demonstrate realtime updates?','Tell audience they exist','Keep role views open and visibly show one action propagating','Refresh everything manually','Show screenshots','B')
q(48,'What if internet access fails during the demo?','Stop presentation','Have cached routes, seeded data and an offline/demo fallback path','Depend entirely on external APIs','Explain what should have happened','B')
q(49,'What if the Driver reports an incorrect checkpoint?','Immediately notify everyone as fact','Keep it pending/unverified until corroborated or reviewed','Delete the Driver','Ignore Driver reports','B')
q(50,'What should happen if Admin rejects a Driver report?','Delete all evidence','Keep report history as rejected; official gate state unchanged; optional reason','Mark gate closed anyway','Alert Trader that route is blocked','B')

# Round 6: offline + route decisions
q(51,'What exactly should offline mode mean in the prototype?','Whole system works fully offline forever','Driver actions are stored locally while disconnected and synchronized later','Only map works offline','Show offline label but do nothing','B')
q(52,'Which Driver actions should be allowed while offline?','Nothing','Status update, checkpoint report, location update, notes and queued photo metadata/files','Admin approval','Trader shipment creation','B')
q(53,'Where should offline Driver actions be stored?','Supabase','IndexedDB/local browser storage until synchronization','React component state only','Python memory','B','B','Explicit final choice: IndexedDB.')
q(54,'What should happen when internet connection returns?','Delete offline records','Sync queued actions in order/independently, mark success, retain failures for retry','Ask Driver to re-enter everything','Upload only most recent action','B')
q(55,'What if a Driver reports a new checkpoint while offline?','Ignore it','Capture GPS/location if available, timestamp, note/photo and temporary local ID in IndexedDB, then sync later','Immediately create official gate','Change route locally for everyone','B','B','GPS can work without internet; if geolocation fails, allow manual pin/location fallback. Demo can simulate coordinates.')
q(56,'What should happen when a newly approved/confirmed checkpoint affects the active route?','Automatically reroute with no Driver control','Mark route disrupted and let Driver choose/reroute operationally','End shipment','Ignore checkpoint','B','B','Final model: system flags disruption/recommendations; Driver can reroute immediately without Admin approval.')
q(57,'Should the Driver be able to choose/suggest an alternate route without Admin approval?','No, never','Yes; Driver owns operational routing and changes are logged/visible','Driver can silently reroute with no record','Trader chooses road','B','B','Final decision explicitly removes Admin approval requirement.')
q(58,'What should happen to the old route after rerouting?','Delete it','Keep it in route history and mark the new route active','Keep both active','Replace every database record','B')
q(59,'What if the Driver reaches a gate that is not on the planned route?','Prevent reporting it','Allow reporting because field reality may differ from planned route data','Automatically reject it','Change destination','B')
q(60,'How complicated should route intelligence become?','Build automatic national route optimization','Stop at simple primary/alternative and Driver rerouting','Build traffic prediction','Add AI route forecasting/recommendation','B','D','User chose D. Final constraint: AI/recommendation engine may suggest/explain routes, but Driver always decides; keep the implementation heuristic and optional after core features.')

# Round 7: RBAC
q(61,'Can a Driver edit an official confirmed gate status?','Yes','No; Driver submits observations, while system/Admin manages authoritative status','Only while driving','Driver and Trader can','B')
q(62,'Can the Driver reroute their own assigned shipment?','No','Yes, without Admin approval, but route change is logged','Only Trader can approve','Only AI decides','B')
q(63,'Can a Driver reroute another Driver’s shipment?','Yes','No','Only if nearby','Only during emergencies','B')
q(64,'Can a Trader see another Trader’s cargo?','Yes','No; only their own shipments','Only on Admin map','Yes but hide cargo names','B')
q(65,'Can a Trader change the Driver’s route?','Yes','No; Trader tracks cargo but does not control field routing','Only if delayed','Always','B')
q(66,'Can Admin see all active trucks?','No','Yes','Only blocked trucks','Only one region','B')
q(67,'Can Admin manually change a Driver’s route?','Never','Yes, as an exceptional operational override with audit reason','Every change must come from Admin','Admin should not see routes','B')
q(68,'What happens if Driver and Admin/system disagree about a gate?','Driver silently overwrites official state','Driver submits a new observation; official state changes only via evidence rules/Admin override','Delete both records','Trader decides','B')
q(69,'Who can see uploaded Driver documents?','Everyone','Relevant Driver + shipment Trader + authorized Admin','Every Trader','Public internet','B')
q(70,'Who should be allowed to mark the shipment DELIVERED?','Trader','Assigned Driver with timestamp/location/proof','Anyone','Only Admin','B')

# Round 8: pheromone concept
q(71,'Should route confidence be global?','Yes, one score for everyone','No; confidence/recommendation can consider context such as vehicle type','Only Admin has score','No score at all','B')
q(72,'What should reinforce a route most?','Driver selecting it','Successfully completing a real traversal','Admin viewing it','Trader liking it','B')
q(73,'Should one Driver travelling a route repeatedly be enough to confirm it?','Yes','No; require multiple unique Drivers','One trip is enough','Admin decides every time','B')
q(74,'What should reduce route confidence?','Nothing','Recent blockages, failed journeys and long inactivity','Driver changing phones','Trader changing cargo','B')
q(75,'Should an unconfirmed route be hidden from other Drivers?','Yes','No; show it clearly labeled unconfirmed','Only Admin sees it','Delete until verified','B')
q(76,'Should the system automatically select the highest-confidence route?','Yes without asking','No; recommend routes and Driver makes final decision','Admin makes final decision','Trader chooses','B')
q(77,'What should a Driver see when choosing among routes?','Only route names','Distance + confidence + recent incidents + last successful use','Only distance','A single AI answer','B')
q(78,'Should route confidence decay over time?','No','Yes','Only manually','Delete routes after one month','B')
q(79,'Should blocked status and confidence be the same number/state?','Yes','No; operational condition and route trust are separate','Only for temporary routes','Remove status entirely','B')
q(80,'How far should you implement the ant-colony idea in this prototype?','Build full academic Ant Colony Optimization','Implement simple pheromone-inspired confidence/recommendation rules','Only mention ants in slides','Replace ORS completely','B')

# Round 9: learned routes + bad data
q(81,'How should a Driver-created route be stored?','One text field','Route record plus ordered GPS breadcrumb points','Screenshot only','Only on Driver phone','B')
q(82,'When should breadcrumb recording start?','Every time app is open','When Driver starts trip or intentionally starts route tracking','Only after delivery','Always in background forever','B')
q(83,'What should count as a successful route traversal?','Driver clicked route','Driver actually travelled most of route and reached intended waypoint/destination','Admin viewed it','Trader received alert','B')
q(84,'How should the system recognize that two Drivers used roughly the same route?','Exact coordinate-for-coordinate match','Compare GPS paths within a reasonable distance tolerance','Same route name only','Ask Admin every time','B')
q(85,'What should happen when a Driver’s GPS data is clearly impossible?','Trust it','Reject/flag impossible jumps and exclude from reinforcement','Confirm route','Delete Driver','B')
q(86,'Should every successful Driver trip add the same amount of confidence?','Yes','Mostly yes for prototype simplicity, but only valid completed trips count','Build complicated ML model','Only Admin trips count','B')
q(87,'How do you stop fake repeated trips from confirming a route?','Do not','Require multiple unique Drivers and valid GPS movement','Trust each Driver forever','Require Trader approval','B')
q(88,'What if Drivers report contradictory conditions on the same route?','Newest report is absolute truth','Preserve observations and mark condition uncertain until more evidence arrives','Delete one','Ask Trader','B')
q(89,'Should Driver reputation affect route confidence?','Build full reputation system','Not in V1; validate evidence rather than ranking people','Only Admin decides reputation','Automatically distrust new Drivers','B')
q(90,'What is the minimum database addition for ant-style route learning?','15 new tables','Add route_points, route_traversals and route metrics','Replace entire database','Store everything in shipments','B')

# Round 10: final DB/API model
q(91,'Should shipments directly contain the whole route history?','Yes, everything in one row','No; shipment stores current route while changes/traversals are separate','Store route history in React','Do not store history','B')
q(92,'What should the final route-related database model be?','routes only','routes + route_points + route_traversals','One table per Driver','GPS data in profiles','B')
q(93,'What should happen when a Driver starts travelling with no known route?','Block Driver','Create a temporary observed-route session and record breadcrumbs','Ask Admin to create road','Fake an ORS route','B')
q(94,'When should an observed path become reusable by other Drivers?','Only after full confirmation','After a valid completed traversal, clearly labeled UNCONFIRMED','Never','After Admin approval only','B')
q(95,'What API pattern should record Driver GPS movement?','One request every second forever','Batch GPS points periodically and support offline queuing','Only upload after delivery','Put coordinates in URL','B')
q(96,'What API should handle unexpected checkpoints?','One generic /update endpoint','Dedicated gate-report endpoints','Direct SQL from React','Put checkpoint data inside route name','B')
q(97,'How should Driver rerouting be recorded?','Just overwrite active_route_id','Create event/traversal history with old/new route, Driver, time and reason','Do not record it','Ask Admin','B')
q(98,'Where should pheromone/confidence calculation happen?','React','FastAPI/backend service','CSS','Driver manually enters score','B')
q(99,'Should confidence be recalculated on every map render?','Yes','No; recalculate on meaningful events and store/update result','Only once','Trader calculates it','B')
q(100,'What should be the single end-to-end backend flow your team gets working first?','AI routing algorithm','Trader shipment → assignment → movement → checkpoint report → alert → reroute → delivery','User profile editing','Analytics dashboard','B')

# Round 11: scope control
q(101,'What is absolutely mandatory for V1?','RBAC, shipment lifecycle, maps, gate reports, alerts, offline simulation/implementation','AI forecasting, analytics, payments, OCR','Full production compliance','Everything discussed so far','A')
q(102,'Should learned/pheromone routes be required for the live demo?','Yes; demo fails without them','No; build only after core logistics flow works','Replace normal routing with them','Remove them completely','B')
q(103,'How much AI route recommendation should V1 contain?','Full ML model','Simple confidence scoring + recommended routes','No explanation','Train a custom LLM','B')
q(104,'Should you implement real background/mobile GPS?','Yes, production mobile GPS plus a simulated fallback','No; simulation only','Build native Android services only','Remove GPS','B','A','Explicit user decision: implement real/live device GPS and simulated GPS fallback. Keep simulation as demo-safe path because assignment only mandates simulated GPS.')
q(105,'How much offline capability should V1 implement?','Complete offline-first application','IndexedDB queue + reconnect sync demonstration','No offline logic','Offline maps for all Myanmar','B')
q(106,'How many shipment scenarios should you prepare?','20 realistic shipments','One perfect golden-path demo + 1–2 seeded examples','Randomly generate hundreds','Only empty screens','B')
q(107,'How many known gates/routes should you seed?','Every gate in Myanmar','Small representative set plus dynamic temporary checkpoints','None','Hundreds of fake gates','B')
q(108,'Should document uploads include OCR/document verification?','Yes','No; upload/display reliably first','Train customs-document model','Remove uploads','B')
q(109,'How much analytics should Admin dashboard contain?','Advanced forecasting/BI','Simple counts/status cards and operational map','No dashboard','Full financial analytics','B')
q(110,'If time runs short, what is the correct feature-cutting order?','Remove core assignment features first','Cut advanced extensions before required features','Keep route AI and remove RBAC','Remove live demo and explain idea','B')

# Round 12: team execution
q(111,'What should the team build first?','Five separate screens in parallel','Shared design package + schema/API contract + one shipment golden path','Slides','Pheromone algorithm','B','CUSTOM','Before B, create V1 architecture diagram, functional diagram, use-case/usage diagram, class diagram and ER diagram in Mermaid/HTML. Then freeze schema/API and build the golden path.')
q(112,'Who should own the shared technical contract?','Everyone changes it freely','Person 1 + Person 3 jointly own it','Frontend only','Nobody','B')
q(113,'What should Person 3 build before frontend developers depend on them?','Every backend feature','Database migrations + seeded users + basic shipment endpoints + API docs','AI routing first','Analytics','B')
q(114,'What should Person 2 build while backend is being prepared?','Wait','Admin/Trader UI against agreed mock JSON','Change schema independently','Build Python','B')
q(115,'What should Person 4 build first?','Full pheromone algorithm','Shared map component + simulated GPS + live GPS mode','Admin dashboard','Slides','B')
q(116,'What should Person 5 test first?','Colors/fonts','Golden-path integration across all three roles','Advanced AI','Presentation animations','B')
q(117,'How should Git workflow work with five people?','Everyone commits to main','Short-lived feature branches + review + protected stable branch/tag','Five separate repos','Send ZIP files','B')
q(118,'When should integration happen?','Final night','Continuously from the first vertical flow','After every feature is finished','During presentation','B')
q(119,'What should happen when AI-generated code causes integration problems?','Keep adding prompts until code grows','Reduce to smallest failing case, inspect contracts, then request targeted fix','Rewrite whole project','Hide bug','B')
q(120,'What should implementation order be?','Maps → AI → UI → DB → auth','Foundation → shipment → maps → gate workflow → offline → learned routes → polish','Pheromone first','Everyone chooses independently','B')

# Round 13: shipment state machine
q(121,'What should the full V1 shipment lifecycle be?','ACTIVE → DONE','REQUESTED → ASSIGNED → PICKED_UP → IN_TRANSIT → AT_CHECKPOINT → CUSTOMS → DELIVERED','25 detailed statuses','Driver types any status','B','B','Later includes CANCELLED as terminal non-delivery state.')
q(122,'Who creates a REQUESTED shipment?','Admin only','Trader','Driver','AI','B')
q(123,'Who changes REQUESTED → ASSIGNED?','Trader','Admin','Driver','Automatically','B')
q(124,'Who changes ASSIGNED → PICKED_UP?','Driver','Trader','Admin only','System timer','A')
q(125,'Who controls normal operational shipment updates after pickup?','Admin','Assigned Driver','Trader','Everyone','B')
q(126,'What if Driver encounters multiple checkpoints?','AT_CHECKPOINT can only happen once','Keep current_status simple and create shipment_event for each checkpoint','Create CHECKPOINT_1/2/3 statuses','Ignore later checkpoints','B')
q(127,'Should the system allow backward/nonlinear state transitions?','Never','Yes when operationally valid, with event history/reason','Any role can change freely','Delete previous events','B')
q(128,'What if Driver skips a state?','Reject every skipped state','Allow reasonable transition and infer/log intermediate event where appropriate','Mark shipment corrupted','Require Admin approval','B')
q(129,'What should happen when shipment is delayed?','Use DELAYED as main lifecycle state','Keep lifecycle status and delay condition separate','Change status to ERROR','Do nothing','B')
q(130,'What exactly should mark a shipment complete?','Driver clicks DELIVERED only','Assigned Driver confirms delivery with timestamp, location and proof photo','Admin closes every shipment','Trader viewing it','B')

# Round 14: diagrams/documentation
q(131,'Should all five diagrams be created before heavy coding?','No; diagrams only for presentation','Yes; create V1 diagrams first and update as implementation changes','Only ER matters','Create after everything is finished','B')
q(132,'What should the Architecture Diagram primarily show?','Database columns','Major technical components and data flow','Every React component','Only users','B')
q(133,'What should the Functional Diagram show?','Classes/methods','Business processes and information flow','SQL relationships','Deployment servers','B')
q(134,'What should the Use-Case Diagram focus on?','Database tables','What Admin, Trader and Driver are allowed to do','CSS components','Python functions','B')
q(135,'Should Use-Case Diagram show Admin approving Driver reroutes?','Yes','No','Only if Trader requests','Always','B')
q(136,'What should the ER Diagram represent?','Python classes','Persistent entities and relationships in Supabase/PostgreSQL','React components','User navigation','B')
q(137,'What should the Class Diagram represent?','Copy ER exactly','Domain objects and application services including behavior','Database indexes','HTML pages only','B')
q(138,'Where should pheromone/route-learning logic appear in diagrams?','Everywhere','Architecture + Class + Functional diagrams, marked as extension','Nowhere','Only ER','B')
q(139,'Should Mermaid files be generated once and forgotten?','Yes','No; keep source under version control and update it','Delete source after PNG','Only Person 1 sees them','B')
q(140,'How should diagrams be used during the 30-minute presentation?','Explain every table/class for 15 minutes','Show only diagrams that support the story; keep details for Q&A','Do not show architecture','Read Mermaid code aloud','B')

# Round 15: map design
q(141,'Should all three roles use the exact same map data?','Yes, everyone sees everything','Same base map component with role-specific layers','Three separate map systems','Only Driver gets map','B')
q(142,'How should different route types appear?','All identical','Visually distinguish planned, active, Driver-observed and disrupted routes','Only active route','Use route names only','B')
q(143,'Should GPS breadcrumbs always be visible?','Yes, every coordinate','No; show current position by default and allow history/trail when useful','Never show breadcrumbs','Only Admin sees location','B')
q(144,'How should an unexpected temporary checkpoint appear?','As confirmed gate immediately','Clearly unverified/pending until confirmed','Hidden until Admin approval','As part of route line','B')
q(145,'What should happen when several reports refer to roughly the same checkpoint?','Show overlapping markers','Combine visually under one checkpoint/incident with multiple observations','Delete all except newest','Hide them','B')
q(146,'How should learned routes display confidence?','Raw decimal only','Simple confidence levels with detailed score on click','Do not show confidence','Only Admin sees confidence','B')
q(147,'What should map do if Driver leaves every known route?','Display error and stop tracking','Continue GPS tracking and begin observed-route trail','Force Driver back','Automatically mark delayed','B')
q(148,'Should ORS routes and Driver-observed routes be treated equally?','Yes','Both are candidates, but source/confidence remain visible','Driver routes always win','ORS always wins','B')
q(149,'How should map clutter be controlled?','Show everything simultaneously','Role-based layers, zoom detail, filters and selected-item focus','Remove information permanently','Separate page for every marker','B')
q(150,'What should happen when Driver taps a route on map?','Immediately switch route','Show route details first, then explicit Use This Route action','Ask Admin','Nothing','B')

# Round 16: alerts
q(151,'Who should receive a Driver’s new checkpoint report immediately?','Everyone','Admin/operations immediately; affected Trader after relevant confirmation','Trader only','Nobody until delivery','B')
q(152,'Should every Driver report create a Trader alert?','Yes','No; only reports/incidents affecting that Trader’s shipment/route','Only if Driver asks','Only once/day','B')
q(153,'What alert severity levels should V1 have?','One generic type','INFO / WARNING / CRITICAL','Ten levels','Free-text severity','B')
q(154,'What should happen when multiple reports describe the same disruption?','Send one alert per report','Update one existing incident/alert rather than spam duplicates','Ignore later reports','Delete original alert','B')
q(155,'Should Admin be able to broadcast a manual alert?','No','Yes; all users, routes or affected shipments/audiences','Only Drivers','Email only','B')
q(156,'What happens when a Driver changes route?','No notification','Log change; Admin and affected Trader see update/lightweight notification','Critical alert every time','Admin must approve first','B')
q(157,'What happens when a route becomes BLOCKED?','Only change marker','Update route/incident state, find affected shipments, create alerts and update maps','Generic message to everyone','Delete route','B')
q(158,'What happens when disruption is resolved?','Leave warning forever','Mark resolved, stop active warning and optionally notify affected users','Delete history','Recreate route','B')
q(159,'Should alerts disappear after being read?','Yes','No; mark read and preserve history','Delete after 30 seconds','Store only unread','B')
q(160,'How should realtime notification delivery work?','Browser refresh only','Store alert in Supabase then realtime subscription updates React; refresh/poll fallback','SMS only','AI chatbot','B')

# Round 17: incident lifecycle
q(161,'What should be the lifecycle of a gate incident?','OPEN → CLOSED only','PENDING → CONFIRMED/UNCERTAIN → RESOLVED → ARCHIVED','Free text only','Delete after resolution','B')
q(162,'What should happen when the first Driver reports a new blockage?','Immediately make official','Create a PENDING incident/report','Ignore it','Alert every user','B')
q(163,'What if a second Driver independently reports the same blockage?','Create completely separate incident','Attach to same incident and increase confidence','Delete first report','Archive both','B')
q(164,'Should multiple matching Driver reports ever auto-confirm an incident?','Never','Yes, optionally after a clear threshold such as 2 unique Drivers in a recent window near same location','One report is enough','Only Trader can confirm','B')
q(165,'What if two Drivers report opposite conditions?','Use whichever came first','Mark incident UNCERTAIN/conflicting and wait for more evidence','Delete both','Trader chooses','B')
q(166,'Who can manually override incident status?','Any Driver','Admin','Trader','Nobody','B')
q(167,'What should happen when Drivers report a previously blocked gate is clear?','Immediately delete blockage','Create CLEAR observations and move toward RESOLVED when enough evidence exists','Ignore them','Keep blocked forever','B')
q(168,'Should resolved incidents still affect route confidence?','No, forget completely','Yes, as historical evidence with reduced weight over time','Permanently punish route','Permanently block route','B')
q(169,'Should Admin have to approve every incident resolution?','Yes, always','No; repeated clear evidence may auto-resolve, Admin handles exceptions','Never allow Admin involvement','Trader approves','B')
q(170,'What should be the source of truth for current route/gate condition?','Latest Driver message only','Aggregated incident state derived from reports plus Admin overrides','Trader comments','OpenRouteService only','B')

# Round 18: route confidence exact heuristic
q(171,'What should starting confidence be for a brand-new Driver-observed route?','0','20','50','100','B')
q(172,'How much should one successful traversal add?','+1','+10','+50','+100','B')
q(173,'Should a successful traversal by a new unique Driver add an extra bonus?','No','Yes, +10 extra','+100 extra','Only Admin counts','B')
q(174,'At what confidence should a route become automatically CONFIRMED?','30','50, and at least 2 unique successful Drivers','100 only','Never automatically','B')
q(175,'How much should a confirmed blockage penalize route confidence?','-2','-20','-100 permanently','No penalty','B')
q(176,'Should an unverified Driver blockage report immediately reduce confidence?','Yes, heavily','Only slightly or not at all until corroborated/confirmed','Set confidence to zero','Delete route','B')
q(177,'How should time decay work?','Reduce every minute','Mild decay after inactivity, e.g. -5 after each 30-day period beyond threshold','Never decay','Delete after 30 days','B')
q(178,'Should confidence ever reach zero?','Yes and delete route','It can become very low, but preserve route history','Never below 50','Always keep 100','B')
q(179,'How should route recommendations combine distance and confidence?','Always choose shortest route','Calculate recommendation score using confidence, incidents and distance','Always highest confidence regardless of distance','Let AI invent answer','B')
q(180,'Should you expose the exact algorithm to Drivers?','Yes, full formula everywhere','No; show understandable reasons, details for Admin/debugging','Hide all reasoning','Show only AI message','B')

# Round 19: live + simulated GPS
q(181,'How should the Driver choose GPS mode?','Automatically with no control','Explicit switch: LIVE GPS or DEMO SIMULATION','Simulation only','Live GPS only','B')
q(182,'Should simulated GPS data affect learned-route confidence?','Yes','No; simulation must never reinforce production routes','Only Admin decides','Sometimes','B')
q(183,'What should happen if Driver denies GPS permission?','App stops','Show warning and allow retry/simulation/manual location fallback','Enable GPS anyway','Log out','B')
q(184,'How often should live GPS points be recorded?','Every millisecond','Controlled interval such as 5–15 seconds or meaningful movement threshold','Once/hour','Only at delivery','B')
q(185,'Should every GPS point immediately be sent to FastAPI?','Yes','No; buffer/batch points, especially for unstable internet','Never send GPS data','Only one final coordinate','B')
q(186,'What should happen when GPS accuracy is poor?','Accept everything','Store accuracy metadata and exclude very unreliable points from route learning','Delete route','Mark delivered','B')
q(187,'How should the simulator move the truck?','Randomly around Myanmar','Along predefined route points with start/pause/next/jump controls','Teleport to destination','Admin manually enters every coordinate','B')
q(188,'Should simulated and live GPS use different map components?','Yes','No; same tracking pipeline, different LocationProvider','Live GPS has no map','Simulation has no map','B')
q(189,'What happens when Driver loses internet but still has GPS?','Stop tracking','Continue recording GPS locally, then synchronize later','Switch to simulation automatically','Delete unsent locations','B')
q(190,'How should route-learning decide whether a traversal is eligible for reinforcement?','Every traversal counts','Require LIVE_GPS + valid movement + successful trip + reasonable path coverage','Simulation counts twice','Driver presses confirm route','B')

# Round 20: documents/photos
q(191,'What kinds of uploads should V1 support?','Any file type','Checkpoint photo, customs/document photo, cargo photo and delivery proof','Video uploads','PDF only','B')
q(192,'Where should uploaded files live?','PostgreSQL binary fields','Supabase Storage with metadata/path in DB','Local phone only','React public folder','B')
q(193,'Should uploaded files be publicly accessible?','Yes','No; private bucket with authorized/signed access','Only checkpoint photos private','Put on social media','B')
q(194,'What metadata should accompany a Driver photo?','File only','Shipment, uploader, type, timestamp, GPS if available and note','Filename only','Password','B')
q(195,'What if Driver takes checkpoint photo while offline?','Reject it','Keep photo locally and queue metadata/upload for synchronization','Upload without internet','Discard image','B')
q(196,'Should system compress large Driver photos before upload?','No','Yes; client-side resizing/compression for reasonable mobile size','Convert to video','Ask Admin','B')
q(197,'Should a Driver be able to hard-delete evidence after uploading it?','Yes permanently','Prefer no hard delete; allow replace/archive while retaining audit metadata','Trader deletes it','Everyone can delete it','B')
q(198,'What should happen when marking shipment DELIVERED?','Status only','Timestamp + GPS + delivery proof photo with optional note','Admin approval every time','Trader types delivered','B')
q(199,'What if delivery-proof upload fails while Driver is offline?','Delivery cannot be recorded','Record pending delivery locally, queue proof upload and sync later','Lose photo','Somehow write to Supabase without internet','B')
q(200,'Should you add OCR/document AI in V1?','Yes, required','No; reliable upload/display first, OCR future work','Build OCR model','Remove uploads','B')

# Round 21: resilience
q(201,'What if OpenRouteService is unavailable during demo?','Stop demo','Fall back to cached/predefined route geometry','Remove map','Retry indefinitely','B')
q(202,'What if Supabase Realtime stops working?','Whole system fails','Keep DB correct and allow polling/manual refresh fallback','Pretend alerts worked','Restart during presentation','B')
q(203,'What if Supabase itself becomes unreachable?','No fallback','Keep seeded local/demo mode for golden-path scenario','Switch databases live','Rewrite backend','B')
q(204,'What if Driver real GPS fails?','Stop tracking','Switch to simulated GPS without changing rest of app','Enter every coordinate manually','Mark delivered','B')
q(205,'What if internet disconnects while Driver is travelling?','Driver cannot do anything','Continue GPS/status/reporting locally through IndexedDB and sync later','Force logout','Switch to Admin','B')
q(206,'What if synchronization partially fails?','Mark everything failed','Sync each queued item independently and retry only failed items','Delete queue','Ask Driver to recreate','B')
q(207,'What if same offline event accidentally syncs twice?','Accept duplicate','Use unique client_event_id/idempotency key','Admin deletes duplicates later','Ignore problem','B')
q(208,'What if an AI-generated code change breaks the stable demo?','Keep fixing directly on main','Maintain known-good demo branch/tag and test changes before merge','Delete Git history','Ask AI to rewrite app','B')
q(209,'What if live demo completely fails despite fallbacks?','End presentation','Have prerecorded golden-path demo as last-resort evidence','Explain verbally for 15 minutes','Show source code','B')
q(210,'What should team test immediately before presenting?','Individual components only','Exact 15-minute golden path from login through delivery, including a failure scenario','Only Admin dashboard','Only pheromone algorithm','B')

# Round 22: presentation
q(211,'How should the 5-minute Problem & Solution section be structured?','Spend all 5 minutes on logistics history','Problem → workflow pain → solution → differentiator → simple architecture/functional visual','Start with DB schema','Explain every feature','B')
q(212,'Should you introduce the pheromone-route concept during first 5 minutes?','Never mention it','Yes, briefly as differentiator without algorithm deep dive','Spend all 5 minutes on ant-colony optimization','Present equations first','B')
q(213,'How should the 15-minute live demo begin?','Admin dashboard','Trader creating a real transport request','Code','Supabase tables','B')
q(214,'What should the 15-minute demo timeline be?','Improvise','Rehearse a tightly timed golden path','Show every screen','Each member demos separate feature','B')
q(215,'Should you intentionally demonstrate offline mode?','No, mention only','Yes; deliberately go offline, create update/report, reconnect and sync','Turn off all presentation internet','Screenshot only','B')
q(216,'How should five people divide speaking?','Everyone exactly 6 minutes','Clear subject ownership with one narrator maintaining continuity','One person 30 minutes','Everyone interrupts when their feature appears','B')
q(217,'What should the 10-minute AI Engineering Reflection focus on?','We used ChatGPT and saved time','Actual AI workflow, prompts, failures, debugging, integration and architecture decisions','Explain LLM math','Show dozens of generated files','B')
q(218,'Should you admit AI-generated failures?','No; make everything perfect','Yes; show 2–3 real meaningful failures and how you diagnosed/fixed them','Blame AI for every bug','Only successful prompts','B')
q(219,'How much architecture should you show during reflection?','Every class/column','One clear architecture diagram; detailed ER/class only when supporting a point or Q&A','No architecture','Read code','B')
q(220,'How should the presentation end?','That is all','Return to Myanmar logistics problem, summarize demonstrated value, state next validation step','End with DB schema','End with AI prompt','B')

# Round 23: AI engineering reflection
q(221,'Should your team save AI prompts during development?','No','Yes; save representative prompts, outputs, failures and corrections','Only screenshots','Recreate later from memory','B')
q(222,'Which AI prompts are most worth saving?','Every conversation','Prompts that produced architecture, schema, UI, debugging and integration decisions','Only perfect prompts','Only short prompts','B')
q(223,'What should a strong coding prompt contain?','Build my logistics app','Context + exact task + constraints + existing schema/API + expected output','Only technology names','Motivational instructions','B')
q(224,'What kind of AI failure should you document?','Only syntax errors','Meaningful failures: invented APIs, wrong schema, security mistakes, integration mismatches','Only funny hallucinations','Do not document failures','B')
q(225,'How should you show a hallucination in presentation?','Read entire conversation','Problem → bad assumption → symptom → correction → lesson','Say AI hallucinated sometimes','Blame tool','B')
q(226,'Should AI-generated code be merged immediately?','Yes','No; review, run, test, fix, then integrate','Only frontend needs testing','AI code assumed correct','B')
q(227,'How should team measure whether AI helped?','Claim 10x faster','Use concrete observations of scaffolding/debugging gains and where manual review remained necessary','No evidence needed','Compare token counts','B')
q(228,'Who should own the AI engineering log?','Nobody','Person 5 collects it; every member contributes examples','Backend only','Presenter invents later','B')
q(229,'Should you show which AI tools each member used?','No','Yes, briefly, explaining what each was useful for','Spend 5 minutes comparing brands','Mention only ChatGPT regardless','B')
q(230,'What should be the main conclusion of AI Engineering Reflection?','AI replaces developers','AI accelerated implementation, but architecture, contracts, verification, integration and testing required human control','Best model built project automatically','Prompt length determines quality','B')

# Round 24: testing
q(231,'What should be the team’s highest-priority test?','Individual component tests only','Full end-to-end shipment lifecycle across all roles','Pheromone score only','UI colors','B')
q(232,'Should RBAC be tested only through UI?','Yes','No; test UI visibility and backend authorization/RLS','Backend only','No tests needed','B')
q(233,'What RBAC tests are mandatory?','Successful login only','Allowed and forbidden actions for all three roles','Admin only','Driver only','B')
q(234,'How should shipment-state transitions be tested?','Normal path only','Normal, repeated, backward-valid and invalid transitions','Let Driver type anything','Only DELIVERED','B')
q(235,'What should be the key offline synchronization test?','Turn off Wi-Fi and stop','Queue several action types, reconnect, verify each syncs exactly once','Queue one text message','Test IndexedDB separately only','B')
q(236,'Should duplicate-sync protection be tested?','No','Yes; send same client_event_id twice and verify one logical event','Delete duplicates manually','Only after deployment','B')
q(237,'How should route-learning tests work?','Use simulation and let it confirm production routes','Use controlled synthetic test data; production reinforcement still requires eligible real GPS traversals','Manual only','Skip tests','B')
q(238,'What should happen in testing when conflicting gate reports arrive?','Latest wins','Verify incident becomes uncertain/conflicted while preserving evidence','Delete both','Block Drivers','B')
q(239,'Who should own final acceptance testing?','Person 5 alone','Person 5 coordinates; all five run golden path at least once','Backend only','Presenter only','B')
q(240,'What should qualify a build as demo-ready?','It compiles','Golden path passes repeatedly, RBAC/fallbacks work, no critical bugs remain','UI looks good','All planned features exist','B')

# Round 25: deployment/security
q(241,'How should frontend be deployed?','Run from developer laptop only','Deploy React/Vite to Vercel with Git-based deployments','Put React in FastAPI templates','Send ZIP','B')
q(242,'How should FastAPI be deployed?','Same browser bundle','Separate backend web service such as Render','Person 3 laptop during presentation','Supabase Storage','B')
q(243,'How many environments should project have?','Production only','Local development + preview/test + stable demo/production','One per person','Ten','B')
q(244,'Which Supabase key should React receive?','Secret key','Publishable key','Database password','Postgres admin credentials','B')
q(245,'Where should Supabase secret key exist?','React public env','FastAPI/server environment only','README','Driver IndexedDB','B')
q(246,'How should environment variables be handled in Git?','Commit .env','Commit .env.example; gitignore real envs; store deployed secrets in hosting dashboards','Paste into JS','Share screenshot','B')
q(247,'Should demo mode use same mutable dataset as daily testing?','Yes','No; maintain predictable seeded demo data that can be reset','Random data','Empty DB before presenting','B')
q(248,'Should feature branches deploy automatically?','No','Yes; preview deployments for integration before stable merge','Every branch replaces production','Localhost only','B')
q(249,'What if a secret key appears in Git or an AI prompt?','Delete line and keep same key','Treat as compromised: rotate/revoke, update envs and redeploy','Rename variable','Hide repository','B')
q(250,'What should final deployment topology be?','Everything on one laptop','Vercel React + Render FastAPI + Supabase + ORS + Device GPS + IndexedDB','Five independent deployments','Supabase only','B')

# Round 26: final DB schema
q(251,'Should profiles duplicate authentication credentials?','Yes, store passwords again','No; Supabase Auth owns auth, profiles stores app identity/role','No profile table','Store passwords in FastAPI','B')
q(252,'What should final shipments table contain?','Every event/GPS point','Current shipment state and references to related entities','Only tracking number','Giant JSON field','B')
q(253,'What should shipment status enums be?','Free text','Fixed lifecycle enum plus separate delay state','One complete boolean','AI invents statuses','B')
q(254,'What belongs in shipment_events?','Only status names','Immutable-ish operational history: status changes, reroutes, checkpoints, delivery etc.','User profiles','Current route geometry','B')
q(255,'What should routes represent?','One specific Driver journey only','Reusable route knowledge','Just ORS URLs','Gate reports','B')
q(256,'How should route geometry and actual journeys be separated?','Put everything in routes','route_points stores geometry; route_traversals stores individual journeys','Screenshots','No geometry history','B')
q(257,'Should gates and incidents be one table?','Yes','No; gates represent locations, incidents represent changing conditions','Only incidents','Only gates','B')
q(258,'What should gate_reports contain?','Only blocked=true','Individual Driver observations linked to gate/incident when possible','Official gate state only','Trader alerts','B')
q(259,'What should final alerts and documents tables look like?','Minimal text only','Structured records linked to shipment/route/incident/user context','React state only','One table for both','B')
q(260,'Should offline-queue records be stored in Supabase as the primary queue?','Yes','No; pending queue lives in Driver IndexedDB, server stores synchronized records and client_event_id','FastAPI RAM only','Nowhere','B')

# Round 27: API contracts
q(261,'How should authentication work between React and FastAPI?','Send username/password to every endpoint','Supabase login → access token → FastAPI validates token and resolves role','FastAPI trusts user_id from React','No auth','B')
q(262,'What should shipment endpoints look like?','One shipment-action endpoint','REST-style endpoints with clear responsibilities','Direct SQL from every screen','One endpoint per button','B')
q(263,'Who should GET /shipments return data for?','Everyone gets everything','Response is restricted by authenticated role','React filters unauthorized records','Only Admin can use it','B')
q(264,'How should Driver location updates work?','One endpoint every second','Batch location endpoint','Manual shipment updates','Coordinates in alerts','B')
q(265,'What endpoint should create checkpoint/gate observations?','/update-route','POST /gate-reports','/admin/gates/change','Put reports inside shipment status','B')
q(266,'How should Admin review reports?','Modify DB directly','Dedicated list/review endpoints for accept/reject/merge and overrides','Delete manually','React sets verified=true directly','B')
q(267,'How should Driver rerouting work?','Admin endpoint only','Assigned Driver calls reroute endpoint with selected route/reason','Trader changes route','React modifies DB row directly','B')
q(268,'What if Driver is following a completely new/unmapped path?','Reroute cannot handle it','Create/finalize an observed route from breadcrumb data','Require ORS route ID','Ask Admin to draw it','B')
q(269,'Where should route confidence logic be exposed?','Driver sends confidence values','Backend calculates and frontend receives score + explanation/reasons','React calculates everything','Admin types score','B')
q(270,'What should API return when something fails?','"error" for everything','Consistent structured error responses with stable error codes','HTML error pages','Empty response','B')

# Round 28: frontend
q(271,'What should happen immediately after login?','Everyone same dashboard','Redirect based on role','User chooses role manually','Always Admin','B')
q(272,'What should shared application shell contain?','Completely different layout for each role','Shared header/navigation shell with role-specific menus','No navigation','Twenty menu items','B')
q(273,'What should Admin navigation contain?','15 sections','Dashboard, Shipments, Operations Map, Incidents/Gates, Broadcast Alerts','Database editor','Dashboard only','B')
q(274,'What should Trader navigation contain?','Everything Admin sees','Dashboard, My Shipments, Alerts','Gate management','Driver management','B')
q(275,'What should Driver navigation look like?','Desktop sidebar with many sections','Mobile-first Current Trip, Map, Report, Sync','Same as Admin','No navigation','B')
q(276,'What components should shared Shipment Detail page have?','One giant component','Reusable ShipmentHeader, map, timeline, route summary, alerts, documents, Driver info','Separate page per status','Plain JSON','B')
q(277,'Should three role maps be separate implementations?','Yes','No; one shared LogisticsMap with role/layer configuration','Only Admin map reusable','Only Driver gets routing','B')
q(278,'What should Driver CurrentTrip screen contain?','Map only','Shipment summary + map + status + key actions + connectivity/sync status','Analytics','Admin approval queue','B')
q(279,'How should alerts appear in UI?','Popups only','Notification badge + persistent alert history + relevant inline warnings','Email only','Map markers only','B')
q(280,'How should frontend loading/error/offline states be handled?','Ignore them','Every important component supports loading, empty, error and offline states','Reload browser on error','Display stack traces','B')

# Round 29: final execution schedule
q(281,'Should all five people start coding different features immediately?','Yes','No; first freeze diagrams, schema, API contract, statuses, RBAC and golden path','Only backend starts','Only frontend starts','B')
q(282,'What should Person 1 own during implementation?','Write most code','Product/integration control, contracts, scope, demo story and unblocking','Maps only','Slides only','B')
q(283,'What is Person 2 implementation sequence?','Polish animations first','Shared UI → Trader → Admin → realtime/alerts → polish','Admin analytics first','Learn FastAPI','B')
q(284,'What is Person 3 implementation sequence?','Pheromone first','Supabase/Auth → schema → FastAPI core → RBAC → incidents → alerts → advanced routing','Deployment first','Presentation first','B')
q(285,'What is Person 4 implementation sequence?','Ant algorithm first','Map → simulation → live GPS → reports → offline → observed routes → recommendations','Slides first','Supabase auth','B')
q(286,'What is Person 5 implementation sequence?','Wait until development finishes','QA from day one + demo data + AI log + presentation assets','Only slides','Only unit tests','B')
q(287,'What should be first shared integration milestone?','Full system','Trader creates shipment → Admin sees/assigns → Driver sees it','AI recommendation','Offline photos','B')
q(288,'What should be second integration milestone?','Presentation slides','Driver movement → map updates → Trader/Admin tracking','Pheromone decay','OCR','B')
q(289,'What should be third integration milestone?','Full route learning','Checkpoint report → incident → alert → reroute → offline synchronization','Analytics','Presentation graphics','B')
q(290,'When should pheromone/learned-route feature be integrated?','Before shipment creation','Only after core handoff, tracking, disruption/offline and delivery are reliable','Never','With authentication','B')

# Round 30: final acceptance
q(291,'When is the project considered functionally complete?','When every planned feature exists','When required golden path works end-to-end and assignment requirements are demonstrated','When pheromone system is perfect','When UI is polished','B')
q(292,'What core assignment features must be demonstrated?','Only shipment tracking','RBAC + three roles + shipment timeline + GPS map + gate/route handling + alerts + offline + Driver photo upload','Pheromone only','Architecture diagrams only','B')
q(293,'What optional feature should be considered the main differentiator?','Dashboard animations','Driver-observed routes + pheromone-inspired route reinforcement','Dark mode','OCR','B')
q(294,'When is documentation considered complete?','README only','Architecture + functional + use-case + ER + class + API + schema + RBAC + testing + AI log + demo script','No docs','Screenshots only','B')
q(295,'What must pass before you freeze demo build?','Application opens','Golden path passes at least three consecutive rehearsals without critical failure','Backend developer says it works','One successful test','B')
q(296,'What fallbacks must be ready before presenting?','None','Cached route + simulated GPS + offline/demo data + realtime refresh fallback + prerecorded emergency demo','Homepage screenshot only','Another laptop with different code','B')
q(297,'What should be finished for AI Engineering Reflection?','List AI tools only','Real examples of prompts, generated output, mistakes, fixes and lessons','Explain LLMs','Say AI saved time','B')
q(298,'What must every one of the five members know before presentation day?','Only personal feature','Overall architecture, golden demo story, speaking part and basic recovery procedure','Every line of code','Slides only','B')
q(299,'When should you stop adding features?','Five minutes before presentation','Once required features are reliable and new changes threaten stability more than value','Never','After implementing all 300 ideas','B')
q(300,'What is the final definition of success?','Maximum feature count','A coherent, working Myanmar-focused prototype demonstrating required workflow, surviving live demo and explaining AI engineering','Best-looking dashboard','Most complicated algorithm','B')

assert len(Q)==300, len(Q)
assert [x['n'] for x in Q] == list(range(1,301))

# --- Utility content ---
assignment_requirements = '''
# Assignment Grounding

The attached assignment asks for a rapid prototype called **Real-Time Logistics Monitoring for Myanmar Trading**. Mandatory demonstrations are:

- Role-based access for **Admin, Trader, Driver**.
- Admin oversees shipments, manages gate status, and broadcasts alerts.
- Trader submits transport requests, tracks their cargo, and receives gate/delay alerts.
- Driver updates location, uploads document photos, and reports status.
- Shipment timeline and simulated live GPS map.
- Myanmar-specific gate/route status behavior, including major-route disruptions.
- Offline capability proposed or simulated for Driver blackouts.
- AI-assisted/vibe-coding methodology with architecture, schema, prompt engineering, integration and debugging.
- 30-minute presentation: **5 min problem/solution + 15 min live demo + 10 min AI engineering reflection**.

**Team-specific adaptation:** the assignment assumes about 11 members, but this project plan is optimized for **5 people**.

**Design extensions beyond the assignment:** real device GPS, Driver-controlled rerouting, dynamic temporary checkpoints, multi-Driver incident consensus, Driver-observed unmapped routes, and pheromone-inspired route confidence/recommendations. These are differentiators, not mandatory assignment requirements.
'''

architecture_md = '''
# System Architecture

## Final stack

- **Frontend:** React + Vite + Tailwind CSS
- **Backend:** Python + FastAPI
- **Identity/Data:** Supabase Auth + PostgreSQL + Realtime + Storage
- **Routing:** OpenRouteService (ORS)
- **Driver location:** Live browser/device GPS + controlled demo simulator
- **Offline:** IndexedDB action/location/file queue
- **Frontend hosting:** Vercel
- **Backend hosting:** Render (or equivalent FastAPI-compatible service)

## Responsibility boundaries

- **React** owns rendering, role-specific navigation, user interaction, connection/offline status, maps and local queue UX.
- **FastAPI** owns privileged business rules, RBAC enforcement, status transitions, incident aggregation, rerouting audit, route scoring and idempotent sync.
- **Supabase** owns authentication, persistence, realtime subscriptions and private file storage.
- **ORS** supplies candidate route geometry and routing; it is not the logistics source of truth.
- **IndexedDB** keeps Driver operations working during blackouts.

## Authority model

- **Admin:** authoritative manual operational oversight and exception handling.
- **Driver:** field evidence + operational route choice for their assigned shipment.
- **Trader:** shipment request and visibility into their own cargo.
- **System consensus:** routine incidents may auto-confirm/resolve from multiple independent Driver reports; Admin can override.
'''

functional_md = '''
# Functional Flows

## Golden shipment flow

1. Trader logs in and creates a transport request.
2. Admin sees the request and assigns a Driver and initial route.
3. Driver starts the trip and chooses Live GPS or Demo Simulation.
4. Location points update the shared logistics map and shipment timeline.
5. Driver encounters a known or unexpected checkpoint.
6. Driver records a gate report with condition, time, location, note and optional photo.
7. If offline, the report/GPS/photo metadata is queued in IndexedDB.
8. On reconnection, queued items synchronize idempotently.
9. The incident engine associates/merges observations; corroborating Drivers can auto-confirm a routine incident; conflicts create UNCERTAIN state; Admin handles overrides/exceptions.
10. Affected Traders receive a relevant alert; Admin and Driver maps update.
11. Driver chooses a route. The system may recommend alternatives, but the Driver makes the operational decision.
12. If the Driver travels through an unmapped path, breadcrumbs become a DRIVER_OBSERVED route.
13. Successful real-GPS traversals reinforce route confidence; diverse Driver evidence can auto-confirm a reusable route.
14. Driver uploads customs/cargo/delivery evidence and marks the shipment delivered.
15. Trader sees delivery state and proof.
'''

rbac_md = '''
# RBAC and Use Cases

## Admin
- View all shipments and trucks.
- Assign Drivers and initial routes.
- View the operations map and all incidents.
- Create known/temporary gates when information comes from non-Driver sources.
- Review/override conflicting or suspicious gate incidents.
- Broadcast alerts to all users, routes or affected shipments.
- Perform exceptional route override with audit reason.

## Trader
- Create shipment request.
- View only own shipments.
- Track route, truck, timeline and alerts.
- View documents/delivery proof relevant to own shipment.
- Cannot change Driver routing or official gate state.

## Driver
- View and update only assigned shipment.
- Start trip and provide live/simulated location.
- Report existing/new checkpoints and field conditions.
- Operate offline and synchronize later.
- Choose/reroute own shipment without routine Admin approval.
- Upload shipment evidence and delivery proof.
- Record/reuse Driver-observed routes.
- Cannot edit authoritative gate state directly or modify another Driver’s shipment.
'''

schema_md = '''
# Database Schema

## Core tables

### profiles
`id`, `full_name`, `role`, `phone`, `is_active`, `created_at`, `updated_at`.
`id` references Supabase Auth user ID.

### shipments
`id`, `tracking_number`, `trader_id`, `driver_id`, `active_route_id`, `cargo_description`, `cargo_type`, `vehicle_type`, origin/destination names and coordinates, `current_status`, `delay_status`, current coordinates, lifecycle timestamps.

### shipment_events
Append-oriented operational history: `shipment_id`, `event_type`, previous/new status, title/description, location, `created_by`, `source`, `client_event_id`, metadata, timestamp.

### routes
Reusable route knowledge: `name`, `source` (ORS/ADMIN/DRIVER_OBSERVED), `trust_status`, `operational_status`, `confidence_score`, endpoints, distance/duration, trip/driver counts, confirmation method, lifecycle timestamps.

### route_points
Ordered geometry points: `route_id`, `sequence`, latitude, longitude.

### route_traversals
Evidence that a Driver actually used a route: route/shipment/driver, location source, start/end, success, reinforcement eligibility, distance, coverage ratio.

### gates
Physical/operational checkpoint locations: name, permanent/temporary type, coordinates, verified flag, created/archived timestamps.

### incidents
Changing conditions associated with a gate and/or route: type, state, severity, confidence, start/resolve times, Admin override flag.

### gate_reports
Individual Driver observations: incident/gate/shipment/driver refs, condition, coordinates, note/photo, verification status, `client_event_id`, report/sync times.

### alerts
User-targeted messages linked to shipment/route/incident: type, severity, title, message, read state/timestamps.

### documents
Private file metadata: shipment/uploader, document type, storage path, MIME/size, capture location/time, note, lifecycle status, optional `client_event_id`.

## Enums
Shipment lifecycle: `REQUESTED`, `ASSIGNED`, `PICKED_UP`, `IN_TRANSIT`, `AT_CHECKPOINT`, `CUSTOMS`, `DELIVERED`, `CANCELLED`.
Delay: `ON_TIME`, `DELAYED`, `UNKNOWN`.
Incident: `PENDING`, `CONFIRMED`, `UNCERTAIN`, `RESOLVED`, `ARCHIVED`.
Route trust: `UNCONFIRMED`, `CONFIRMED`.
Route operation: `CLEAR`, `DISRUPTED`, `BLOCKED`, `UNKNOWN`.
'''

api_md = '''
# FastAPI Contract

## Authentication
React signs in via Supabase Auth and sends `Authorization: Bearer <access-token>`. FastAPI validates the token and resolves user ID/role. Backend must never trust a role or user ID merely because the browser submitted it.

## Shipments
- `POST /shipments`
- `GET /shipments`
- `GET /shipments/{id}`
- `POST /shipments/{id}/assign`
- `POST /shipments/{id}/status`
- `POST /shipments/{id}/reroute`
- `POST /shipments/{id}/deliver`
- `POST /shipments/{id}/locations/batch`

## Gate reports / incidents
- `POST /gate-reports`
- `GET /gate-reports?status=PENDING`
- `POST /gate-reports/{id}/accept`
- `POST /gate-reports/{id}/reject`
- `POST /gate-reports/{id}/merge`

## Routes
- `GET /routes`
- `GET /routes/recommendations?shipment_id=...`
- `POST /routes/observed`

## Documents
- `POST /documents`
- `GET /shipments/{id}/documents`

## Alerts
- `GET /alerts`
- `POST /alerts/{id}/read`
- `POST /admin/alerts/broadcast`

## Error envelope
```json
{
  "error": {
    "code": "SHIPMENT_NOT_ASSIGNED_TO_DRIVER",
    "message": "You cannot update this shipment.",
    "details": null
  }
}
```
Stable codes include `FORBIDDEN`, `NOT_FOUND`, `INVALID_STATUS_TRANSITION`, `DUPLICATE_EVENT`, `GPS_VALIDATION_FAILED`, `ROUTE_NOT_AVAILABLE`, `OFFLINE_SYNC_CONFLICT`.
'''

frontend_md = '''
# Frontend Specification

## Shared components
`AppShell`, `Header`, `UserMenu`, `ConnectionStatus`, `NotificationBell`, `LogisticsMap`, `ShipmentHeader`, `ShipmentStatusBadge`, `ShipmentTimeline`, `RouteSummary`, `AlertPanel`, `DocumentGallery`, `DriverInfo`.

## Admin routes
`/admin`, `/admin/shipments`, `/admin/map`, `/admin/incidents`, `/admin/alerts`.
Admin map shows all active trucks, known/temporary gates, incidents and routes.

## Trader routes
`/trader`, `/trader/shipments`, `/trader/shipments/new`, `/trader/shipments/:id`, `/trader/alerts`.
Trader sees only their shipments, current route/truck, timeline, relevant incidents, documents and delivery proof.

## Driver routes
`/driver`, `/driver/map`, `/driver/report`, `/driver/routes`, `/driver/sync`.
Mobile-first CurrentTrip screen shows shipment, map, status, Update Status, Report Checkpoint, Choose Route, Upload Document, Mark Delivered, connection state and pending sync count.

## Map layers
One `LogisticsMap` implementation with role-based visibility. Planned/active/observed/disrupted route types are visually distinguishable. Current truck is default; breadcrumb history is optional. Temporary checkpoints display pending/unverified state until corroborated/confirmed.
'''

offline_md = '''
# Offline, GPS and Mapping

## Location provider abstraction
A shared `LocationProvider` interface is implemented by `LiveGPSProvider` and `SimulationProvider`. The tracking and map layers should not care which provider is active.

## Live GPS
Capture latitude, longitude, accuracy, timestamp and source. Sample every 5–15 seconds or after meaningful movement. Reject impossible jumps and ignore low-quality points for route reinforcement.

## Simulation
Use predefined route points with start/pause/next/jump-to-checkpoint controls. Simulation demonstrates behavior but never reinforces production route confidence.

## IndexedDB queue
Queue types: GPS batch, status update, gate report, document upload, delivery event. Queue states: PENDING, SYNCING, SYNCED, FAILED. Every action gets `client_event_id` for idempotent synchronization. Partial failure must not block independent queue items.
'''

route_md = '''
# Driver-Observed Routes and Pheromone-Inspired Confidence

This is a team extension, not an assignment requirement.

## Concept
A Driver may leave the known/planned route. The app continues recording GPS breadcrumbs and stores a Driver-observed route rather than forcing the Driver back onto mapped roads. Other Drivers may reuse an unconfirmed route with a clear warning.

## Simplified reinforcement heuristic
- New Driver-observed route starts at confidence **20**.
- Valid successful traversal: **+10**.
- First successful traversal by a new unique Driver: **+10 diversity bonus**.
- Confirmed blockage: **-20** confidence (operational BLOCKED state remains separate).
- Inactivity beyond 30 days: mild **-5 decay** per configured period.
- Auto-confirm when confidence ≥ **50** and at least **2 unique Drivers** have successfully used the route.

Only LIVE_GPS traversals that pass validation and have reasonable path coverage are eligible for reinforcement. The recommendation engine considers confidence, active incidents and distance, but the Driver makes the route decision.
'''

incidents_md = '''
# Gates, Incidents and Alerts

## Separation of concepts
- **Gate:** a location/checkpoint.
- **Gate report:** one Driver observation.
- **Incident:** the system’s aggregated current interpretation of conditions.
- **Alert:** a user-targeted message generated because an incident matters to them.

## Incident logic
First observation creates PENDING state. Multiple matching reports from unique Drivers in a recent location/time window can auto-confirm. Opposing reports produce UNCERTAIN state. Repeated CLEAR evidence can auto-resolve. Admin handles overrides, suspicious data and exception cases.

## Alerts
Severities: INFO, WARNING, CRITICAL. Avoid spam by updating one evolving incident/alert rather than generating duplicate warnings for every report. Traders receive only alerts relevant to their shipment/route. Admin can broadcast manually by audience.
'''

deploy_md = '''
# Security and Deployment

## Topology
React/Vite on Vercel → HTTPS → FastAPI on Render → Supabase Auth/Postgres/Realtime/Storage. ORS and device GPS are external inputs; IndexedDB provides Driver offline continuity.

## Secrets
Browser receives only the Supabase **publishable** key. FastAPI holds the Supabase **secret** key and ORS server-side credentials where applicable. Commit `.env.example`; never commit real `.env` files. If a key appears in Git or an AI prompt, treat it as compromised and rotate/revoke it.

## Environments
Local development, preview/test deployments, and stable demo/production. Maintain predictable seeded demo data and a known-good tagged build.
'''

testing_md = '''
# Testing and Demo Readiness

## Highest-priority test
Run the complete golden path across Trader, Admin and Driver.

## RBAC negative tests
Trader cannot view another Trader’s shipment or approve incidents. Driver cannot update another Driver’s shipment. Admin can view all shipments and perform authorized operations.

## Offline test
Queue GPS batch + status + gate report + photo while offline, reconnect, verify each logical action appears exactly once. Retry only failed items.

## Route-learning test
Synthetic test data may exercise the scoring function, but app-level production reinforcement must reject simulated traversals.

## Demo-ready gate
Three consecutive golden-path rehearsals, no critical bugs, RBAC negative tests pass, cached-route and simulation fallbacks work, and the emergency recording is available.
'''

ai_md = '''
# AI Engineering Reflection

Maintain `/docs/ai-log/` throughout development. Each member contributes at least one useful prompt, one AI-generated mistake, the correction and the lesson.

## Entry template
- **Task:** What were we trying to do?
- **Prompt:** Exact or representative prompt.
- **Output:** What the model generated.
- **Problem:** What was wrong or incomplete?
- **Fix:** What humans changed.
- **Lesson:** How prompt/context/review should improve.

Strong examples include schema-field hallucinations, inconsistent API contracts, incorrect RBAC assumptions, exposed credentials, offline assumptions, duplicated business logic and integration mismatches.

Core reflection: AI accelerates scaffolding, debugging and implementation, but architecture, shared contracts, security, verification, integration and testing remain human responsibilities.
'''

demo_md = '''
# Presentation and Demo Script

## Required 30-minute structure
- **0–5 min:** Problem & Solution.
- **5–20 min:** Live Demo.
- **20–30 min:** AI Engineering Reflection.

## Suggested 15-minute live flow
1. Trader login and create SHP-001.
2. Admin login; assign Driver and route.
3. Driver login; start trip and GPS tracking.
4. Deliberately enter offline mode.
5. Driver discovers/reports unexpected checkpoint offline.
6. Reconnect; queue synchronizes.
7. Incident becomes confirmed through Admin/consensus; Trader alert appears.
8. Driver chooses alternate/unmapped route; maps update.
9. Driver uploads document/evidence.
10. Driver marks delivered with proof; Trader sees delivery.

## Fallback ladder
1. Live system.
2. Demo/fallback mode with cached route/simulated GPS/seeded data.
3. Prerecorded golden-path demo as emergency last resort.
'''

execution_md = '''
# Five-Person Execution Plan

## Person 1 - Product + Integration Lead
Owns requirements, diagrams, contracts, scope, Git integration, demo narrative and unblock decisions.

## Person 2 - React Admin + Trader
Shared UI → Trader flow → Admin flow → realtime/alerts → polish. Develop against agreed mock JSON while backend is incomplete.

## Person 3 - FastAPI + Supabase
Supabase/Auth → migrations/seeds → token validation/RBAC → shipment API → events → gate/incidents → alerts/documents → route confidence → deployment.

## Person 4 - Driver + Map + GPS + Offline
Shared map → cached route → simulator → ORS → live GPS → Driver actions → checkpoint reporting → IndexedDB → observed routes → recommendations.

## Person 5 - QA + Demo + AI Reflection
Test plan from day one, RBAC/integration tests, seeded demo data, AI log collection, demo-readiness checklist, slides and fallback recording.

## Integration milestones
1. Trader creates → Admin assigns → Driver sees.
2. Driver movement → Admin/Trader map tracking.
3. Checkpoint report → incident → alert → reroute.
4. Offline queue → reconnect sync.
5. Documents + delivery proof.
6. Learned routes only after milestones 1–5 are stable.
'''

# Mermaid files
mermaids = {
'architecture.mmd': r'''flowchart LR
  A[Admin Web] --> FE[React + Vite + Tailwind]
  T[Trader Web] --> FE
  D[Driver Mobile Web] --> FE
  FE -->|Bearer token / HTTPS| API[FastAPI]
  FE <-->|Auth / Realtime safe client ops| SB[Supabase]
  API --> SB
  SB --> DB[(PostgreSQL)]
  SB --> AUTH[Auth]
  SB --> STORE[Private Storage]
  SB --> RT[Realtime]
  API --> ORS[OpenRouteService]
  D --> GPS[Device GPS]
  D --> IDB[(IndexedDB Offline Queue)]
  IDB -->|Reconnect sync| API
  API --> RCE[Route Confidence Engine]
''',
'functional-flow.mmd': r'''flowchart TD
  T[Trader creates shipment] --> A[Admin assigns Driver + initial route]
  A --> S[Driver starts trip]
  S --> G[GPS updates map + timeline]
  G --> C{Unexpected checkpoint?}
  C -- No --> G
  C -- Yes --> O{Online?}
  O -- Yes --> R[Submit gate report]
  O -- No --> I[Queue in IndexedDB]
  I --> X[Reconnect + idempotent sync]
  X --> R
  R --> E[Incident engine aggregates evidence]
  E --> V{Confirmed / uncertain?}
  V --> AL[Relevant alerts + map updates]
  AL --> RR[Driver chooses route]
  RR --> U{Known route?}
  U -- Yes --> G
  U -- No --> OBS[Record Driver-observed breadcrumbs]
  OBS --> G
  G --> DEL[Driver uploads proof + marks Delivered]
''',
'use-case.mmd': r'''flowchart LR
  Admin((Admin)) --> A1[View all shipments]
  Admin --> A2[Assign Driver]
  Admin --> A3[View operations map]
  Admin --> A4[Manage/override incidents]
  Admin --> A5[Broadcast alerts]
  Trader((Trader)) --> T1[Create shipment request]
  Trader --> T2[Track own shipment]
  Trader --> T3[View timeline/map]
  Trader --> T4[Receive alerts]
  Trader --> T5[View delivery proof]
  Driver((Driver)) --> D1[View assigned trip]
  Driver --> D2[Update status/location]
  Driver --> D3[Report checkpoint]
  Driver --> D4[Choose/reroute]
  Driver --> D5[Upload documents]
  Driver --> D6[Operate offline + sync]
  Driver --> D7[Record observed route]
  Driver --> D8[Mark delivered]
''',
'er.mmd': r'''erDiagram
  PROFILES ||--o{ SHIPMENTS : "trader/driver"
  PROFILES ||--o{ SHIPMENT_EVENTS : creates
  PROFILES ||--o{ GATE_REPORTS : reports
  PROFILES ||--o{ ALERTS : receives
  PROFILES ||--o{ DOCUMENTS : uploads
  SHIPMENTS ||--o{ SHIPMENT_EVENTS : has
  SHIPMENTS ||--o{ DOCUMENTS : has
  SHIPMENTS ||--o{ GATE_REPORTS : context
  SHIPMENTS ||--o{ ROUTE_TRAVERSALS : has
  ROUTES ||--o{ ROUTE_POINTS : geometry
  ROUTES ||--o{ ROUTE_TRAVERSALS : evidence
  ROUTES ||--o{ INCIDENTS : affected_by
  GATES ||--o{ INCIDENTS : has
  GATES ||--o{ GATE_REPORTS : observed_at
  INCIDENTS ||--o{ GATE_REPORTS : aggregates
  INCIDENTS ||--o{ ALERTS : triggers
''',
'classes.mmd': r'''classDiagram
  class Shipment {
    +UUID id
    +ShipmentStatus currentStatus
    +DelayStatus delayStatus
    +UUID activeRouteId
    +assignDriver()
    +updateStatus()
    +reroute()
    +markDelivered()
  }
  class Route {
    +RouteSource source
    +TrustStatus trustStatus
    +OperationalStatus operationalStatus
    +float confidenceScore
    +reinforce()
    +applyDecay()
    +markDisrupted()
  }
  class GateReport {
    +Condition reportedCondition
    +VerificationStatus verificationStatus
    +accept()
    +reject()
    +merge()
  }
  class Incident {
    +IncidentState state
    +Severity severity
    +float confidence
    +aggregateEvidence()
    +resolve()
  }
  class ShipmentService
  class GateIncidentService
  class RoutingService
  class OfflineSyncService
  class AlertService
  class RouteConfidenceService
  ShipmentService --> Shipment
  RoutingService --> Route
  GateIncidentService --> GateReport
  GateIncidentService --> Incident
  RouteConfidenceService --> Route
''',
'shipment-state-machine.mmd': r'''stateDiagram-v2
  [*] --> REQUESTED
  REQUESTED --> ASSIGNED
  REQUESTED --> CANCELLED
  ASSIGNED --> PICKED_UP
  ASSIGNED --> CANCELLED
  PICKED_UP --> IN_TRANSIT
  IN_TRANSIT --> AT_CHECKPOINT
  AT_CHECKPOINT --> IN_TRANSIT
  IN_TRANSIT --> CUSTOMS
  CUSTOMS --> IN_TRANSIT
  CUSTOMS --> DELIVERED
  IN_TRANSIT --> DELIVERED
  DELIVERED --> [*]
  CANCELLED --> [*]
''',
'incident-lifecycle.mmd': r'''stateDiagram-v2
  [*] --> PENDING
  PENDING --> CONFIRMED: corroborated/manual confirm
  PENDING --> UNCERTAIN: conflicting evidence
  UNCERTAIN --> CONFIRMED: stronger corroboration
  CONFIRMED --> RESOLVED: repeated CLEAR/manual resolution
  UNCERTAIN --> RESOLVED: evidence clears issue
  RESOLVED --> ARCHIVED
  ARCHIVED --> [*]
''',
'offline-sync.mmd': r'''flowchart TD
  A[Driver action / GPS / photo] --> N{Online?}
  N -- Yes --> API[FastAPI]
  N -- No --> IDB[(IndexedDB)]
  IDB --> Q[Queue item with client_event_id]
  Q --> C{Connection restored?}
  C -- No --> Q
  C -- Yes --> API
  API --> V[Validate + authorize + deduplicate]
  V --> DB[(Supabase)]
  DB --> OK[Mark local item SYNCED]
  V -->|item fails| F[Mark FAILED; retry independently]
''',
'route-learning.mmd': r'''flowchart TD
  A[Driver leaves known route] --> B[Record live GPS breadcrumbs]
  B --> C[Complete traversal]
  C --> D{Valid LIVE_GPS + coverage?}
  D -- No --> X[Store history only]
  D -- Yes --> R[Create/reuse DRIVER_OBSERVED route]
  R --> S[+10 successful traversal]
  S --> U{New unique Driver?}
  U -- Yes --> V[+10 diversity bonus]
  U -- No --> W[No diversity bonus]
  V --> T{confidence >= 50 and >=2 unique Drivers?}
  W --> T
  T -- Yes --> CF[CONFIRMED]
  T -- No --> UC[UNCONFIRMED]
  CF --> P[Recommendation engine]
  UC --> P
  P --> DRI[Driver decides]
''',
'deployment.mmd': r'''flowchart LR
  GH[Git Repository] --> V[Vercel: React/Vite]
  GH --> R[Render: FastAPI]
  V -->|HTTPS| R
  V -->|Auth / Realtime| S[Supabase]
  R --> S
  S --> DB[(PostgreSQL)]
  S --> ST[Private Storage]
  S --> AU[Auth]
  R --> ORS[OpenRouteService]
  DR[Driver Device] --> GPS[Live GPS]
  DR --> IDB[(IndexedDB)]
  IDB --> R
''',
'golden-path-sequence.mmd': r'''sequenceDiagram
  participant T as Trader
  participant A as Admin
  participant D as Driver
  participant UI as React
  participant API as FastAPI
  participant S as Supabase
  T->>UI: Create shipment
  UI->>API: POST /shipments
  API->>S: Persist REQUESTED
  A->>UI: Assign Driver
  UI->>API: POST /shipments/{id}/assign
  API->>S: ASSIGNED + event
  D->>UI: Start trip / GPS
  UI->>API: Location batches
  D->>UI: Offline checkpoint report
  UI-->>UI: Queue in IndexedDB
  D->>UI: Reconnect
  UI->>API: Sync gate report
  API->>S: Report + incident
  API->>S: Alert affected Trader
  D->>UI: Choose reroute
  UI->>API: POST /shipments/{id}/reroute
  API->>S: Route change event
  D->>UI: Upload proof + Deliver
  UI->>API: POST /shipments/{id}/deliver
  API->>S: DELIVERED + proof
  S-->>T: Realtime/refresh update
'''
}

# Simple HTML wrapper for Mermaid source
html_template='''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>body{{font-family:Arial,sans-serif;margin:0;padding:24px;background:#f7f7f8;color:#161616}}main{{max-width:1200px;margin:auto;background:white;padding:24px;border-radius:12px;box-shadow:0 1px 8px #0001}}pre.mermaid{{overflow:auto}}</style></head><body><main><h1>{title}</h1><pre class="mermaid">{code}</pre></main><script type="module">import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs';mermaid.initialize({{startOnLoad:true,securityLevel:'loose'}});</script></body></html>'''
for name,code in mermaids.items():
    (BASE/'mermaid'/name).write_text(code.strip()+"\n")
    title=name.replace('.mmd','').replace('-',' ').title()
    (BASE/'html'/(name.replace('.mmd','.html'))).write_text(html_template.format(title=title,code=code))

# Decision log Markdown
md=['# Decision Log: Grill-Me Questions 001–300','', 'This file preserves every grill question, its four choices, the recommendation, the final team decision, and later refinements where applicable.','', assignment_requirements]
for x in Q:
    if (x['n']-1)%10==0:
        md += [f"\n# Round {(x['n']-1)//10+1}: Questions {x['n']}–{x['n']+9}", '']
    md += [f"## Q{x['n']}. {x['question']}",
           f"- **A.** {x['choices']['A']}", f"- **B.** {x['choices']['B']}", f"- **C.** {x['choices']['C']}", f"- **D.** {x['choices']['D']}",
           f"- **Original recommendation:** {x['recommendation']}", f"- **Final decision:** {x['final']}"]
    if x['note']:
        md += [f"- **Refinement / rationale:** {x['note']}"]
    else:
        md += ["- **Refinement / rationale:** Accepted as recommended. The final system specification and thematic documents in this pack define the implementation consequence."]
    md += ['']
(BASE/'decision-log-001-300.md').write_text('\n'.join(md))

# Topical docs
files = {
'01-assignment-grounding.md': assignment_requirements,
'02-system-architecture.md': architecture_md,
'03-functional-flows.md': functional_md,
'04-use-cases-rbac.md': rbac_md,
'05-database-schema.md': schema_md,
'06-api-contract.md': api_md,
'07-frontend-specification.md': frontend_md,
'08-offline-gps-mapping.md': offline_md,
'09-route-learning-pheromone.md': route_md,
'10-incidents-alerts.md': incidents_md,
'11-security-deployment.md': deploy_md,
'12-testing-demo-readiness.md': testing_md,
'13-ai-engineering-reflection.md': ai_md,
'14-demo-script-presentation.md': demo_md,
'15-five-person-execution-plan.md': execution_md,
}
for fn,content in files.items(): (BASE/fn).write_text(content.strip()+"\n")
(BASE/'testing'/'demo-readiness.md').write_text(testing_md.strip()+"\n")
(BASE/'ai-log'/'reflection-summary.md').write_text(ai_md.strip()+"\n")
(BASE/'README.md').write_text('''# Myanmar Logistics & Tracking Prototype Documentation Pack\n\nThis pack consolidates the attached assignment and all 300 Grill-Me design decisions into one implementation source of truth for a five-person team.\n\nStart with:\n1. `decision-log-001-300.md` for every question and decision.\n2. `02-system-architecture.md` through `15-five-person-execution-plan.md` for implementation specifications.\n3. `mermaid/` for editable Mermaid source diagrams.\n4. `html/` for browser-viewable Mermaid diagram pages (requires internet to load Mermaid JS).\n5. `Myanmar_Logistics_Project_Master_Documentation.docx` for the consolidated Word document.\n\nThe pack clearly distinguishes assignment requirements from team extensions such as real GPS, Driver-controlled rerouting and pheromone-inspired learned routes.\n''')
shutil.copy('/mnt/data/README-From-assignment.txt', BASE/'assignment-source.txt')

# --- DOCX creation ---
doc = Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.65); sec.bottom_margin=Inches(0.65); sec.left_margin=Inches(0.7); sec.right_margin=Inches(0.7)

# styles
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); styles['Normal'].font.size=Pt(9.2)
for st,size,color in [('Title',26,'17365D'),('Heading 1',17,'17365D'),('Heading 2',13,'2F5597'),('Heading 3',11,'44546A')]:
    s=styles[st]; s.font.name='Aptos Display' if st!='Normal' else 'Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'),s.font.name); s._element.rPr.rFonts.set(qn('w:hAnsi'),s.font.name); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)

# custom compact styles
if 'Decision' not in styles:
    s=styles.add_style('Decision',WD_STYLE_TYPE.PARAGRAPH); s.font.name='Aptos'; s.font.size=Pt(9.2); s.font.bold=True; s.font.color.rgb=RGBColor(31,78,121)
if 'Choice' not in styles:
    s=styles.add_style('Choice',WD_STYLE_TYPE.PARAGRAPH); s.font.name='Aptos'; s.font.size=Pt(8.8)

# title page
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Myanmar Logistics & Tracking System\n'); r.bold=True; r.font.size=Pt(28); r.font.color.rgb=RGBColor(23,54,93)
r=p.add_run('Complete Project Documentation & 300-Decision Log'); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(47,85,151)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Five-Person Rapid Prototype Plan\nReact + FastAPI + Supabase + OpenRouteService + IndexedDB'); r.font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_before=Pt(24)
r=p.add_run('Source basis: attached assignment + all Grill-Me decisions in this conversation.'); r.italic=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Assignment deliverable: Live Prototype Demo & 30-Minute Presentation, 20 Sept 2026.'); r.bold=True

doc.add_page_break()

# Executive summary
h=doc.add_heading('Executive Summary',1)
for para in [
    'This document is the consolidated source of truth for the five-person Myanmar logistics prototype. It records every design question from Q001 to Q300, including user overrides, later refinements and the final implementation consequence.',
    'The core product is a role-based logistics monitoring system: Traders request and track cargo, Drivers update field conditions and operational routes, and Admins oversee shipments and authoritative incident state. The system is designed for unstable connectivity and changing road/checkpoint conditions.',
    'Mandatory assignment features remain the delivery gate: three roles, shipment timeline, simulated live GPS, route/gate disruption handling, Trader alerts, Driver document-photo upload, offline capability, and AI-engineering reflection. Real GPS and pheromone-inspired learned routes are differentiators, not prerequisites.'
]: doc.add_paragraph(para)

# Quick TOC-ish list
for title in ['Part I - Assignment and Final Product Specification','Part II - Architecture and Technical Documentation','Part III - Five-Person Execution, Testing and Presentation','Part IV - Complete Decision Log Q001-Q300','Appendix - Mermaid Diagram Index']:
    doc.add_paragraph(title, style='Heading 2')

def add_md_section(title, text):
    doc.add_page_break(); doc.add_heading(title,1)
    for line in text.strip().splitlines():
        line=line.rstrip()
        if not line: continue
        if line.startswith('# '):
            # skip duplicate title
            continue
        elif line.startswith('## '): doc.add_heading(line[3:],2)
        elif line.startswith('### '): doc.add_heading(line[4:],3)
        elif line.startswith('- '): doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith('```'):
            continue
        else:
            p=doc.add_paragraph(line)
            if line.startswith('**') and line.endswith('**'): p.runs[0].bold=True

# Part I-III sections
add_md_section('Part I - Assignment Grounding', assignment_requirements)
add_md_section('Final System Architecture', architecture_md)
add_md_section('Functional Flows', functional_md)
add_md_section('Use Cases and RBAC', rbac_md)
add_md_section('Database Schema', schema_md)
add_md_section('FastAPI Contract', api_md)
add_md_section('Frontend Specification', frontend_md)
add_md_section('Offline, GPS and Mapping', offline_md)
add_md_section('Driver-Observed Routes and Pheromone-Inspired Confidence', route_md)
add_md_section('Gates, Incidents and Alerts', incidents_md)
add_md_section('Security and Deployment', deploy_md)
add_md_section('Testing and Demo Readiness', testing_md)
add_md_section('AI Engineering Reflection', ai_md)
add_md_section('Presentation and Demo Script', demo_md)
add_md_section('Five-Person Execution Plan', execution_md)

# Decision log in detail

doc.add_page_break(); doc.add_heading('Part IV - Complete Decision Log Q001-Q300',1)
doc.add_paragraph('Every Grill-Me question is reproduced below with four choices, original recommendation, final decision, and any later refinement. When the final decision says CUSTOM, the refinement text is authoritative.')
for idx,x in enumerate(Q):
    if idx%10==0:
        if idx>0: doc.add_page_break()
        rnd=idx//10+1
        doc.add_heading(f'Round {rnd}: Questions {x["n"]:03d}-{x["n"]+9:03d}',1)
    doc.add_heading(f'Q{x["n"]:03d}. {x["question"]}',2)
    for letter in 'ABCD':
        p=doc.add_paragraph(style='Choice');
        rr=p.add_run(f'{letter}. '); rr.bold=True
        p.add_run(x['choices'][letter])
    p=doc.add_paragraph(style='Decision'); p.add_run(f'Original recommendation: {x["recommendation"]}')
    p=doc.add_paragraph(style='Decision'); p.add_run(f'Final decision: {x["final"]}')
    p=doc.add_paragraph(); rr=p.add_run('Refinement / implementation consequence: '); rr.bold=True
    p.add_run(x['note'] if x['note'] else 'Accepted as recommended. See the relevant technical section in this document for the detailed implementation consequence.')

# Mermaid appendix
doc.add_page_break(); doc.add_heading('Appendix - Mermaid Diagram Index',1)
doc.add_paragraph('Editable Mermaid source is included in the documentation pack under /mermaid. Browser-viewable HTML wrappers are under /html.')
for name in mermaids:
    doc.add_heading(name,2)
    code=mermaids[name].strip()
    for line in code.splitlines():
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25)
        r=p.add_run(line); r.font.name='Liberation Mono'; r._element.rPr.rFonts.set(qn('w:ascii'),'Liberation Mono'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Liberation Mono'); r.font.size=Pt(7.3)

# Header/footer
for section in doc.sections:
    hp=section.header.paragraphs[0]; hp.text='Myanmar Logistics & Tracking System - Master Documentation'; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.size=Pt(8); hp.runs[0].font.color.rgb=RGBColor(100,100,100)
    fp=section.footer.paragraphs[0]; fp.text='Five-Person Prototype | Source: attached assignment + Grill-Me decisions'; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fp.runs[0].font.size=Pt(8); fp.runs[0].font.color.rgb=RGBColor(120,120,120)

DOCX=BASE/'Myanmar_Logistics_Project_Master_Documentation.docx'
doc.save(DOCX)

# machine-readable decisions too
(BASE/'decision-log-001-300.json').write_text(json.dumps(Q,ensure_ascii=False,indent=2))

# ZIP everything except generator itself and rendered QA later
ZIP=Path('/mnt/data/Myanmar_Logistics_Documentation_Pack.zip')
with zipfile.ZipFile(ZIP,'w',zipfile.ZIP_DEFLATED) as z:
    for p in BASE.rglob('*'):
        if p.is_file() and p.name!='generate_pack.py':
            z.write(p, p.relative_to(BASE))
print(DOCX)
print(ZIP)
